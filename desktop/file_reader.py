import csv
import json
from pathlib import Path

from deduplication.normalization import normalize_name
from paths import data_dir
from dataclasses import dataclass

SYNONYMS = {
    "nom": {"nom", "nom de famille", "last name", "lastname", "surname", "family name"},
    "prenom": {"prenom", "first name", "firstname", "given name"},
    "email": {"email", "e mail", "mail", "adresse email", "adresse mail",
              "adresse e mail", "courriel"},
    "telephone": {"telephone", "tel", "phone", "phone number", "mobile",
                  "numero de telephone", "contact"},
    "sexe": {"sexe", "genre", "gender", "sex"},
    "formation": {"formation", "training", "training name", "name of training conducted",
                  "name of training", "nom de la formation", "intitule de la formation"},    
}

_STORE = data_dir() / "column_mappings.json"


def guess_mapping(headers) -> dict[str, str]:
    """Associe chaque champ à la colonne qui porte un nom courant (nom, e-mail, tél...)."""
    mapping: dict[str, str] = {}
    ambiguous = None
    for h in headers:
        key = normalize_name(h)
        if key == "name":
            ambiguous = ambiguous or h
        for field, names in SYNONYMS.items():
            if field not in mapping and key in names:
                mapping[field] = h
    # « Name » seul est ambigu : on le range dans le champ identité qui manque
    if ambiguous:
        if "nom" in mapping and "prenom" not in mapping:
            mapping["prenom"] = ambiguous
        elif "prenom" in mapping and "nom" not in mapping:
            mapping["nom"] = ambiguous
    return mapping


def _signature(headers) -> str:
    return "|".join(headers)


def load_saved_mapping(headers) -> dict[str, str] | None:
    try:
        data = json.loads(_STORE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    saved = data.get(_signature(headers))
    if isinstance(saved, dict) and all(v in headers for v in saved.values()):
        return saved
    return None


def save_mapping(headers, mapping: dict[str, str]) -> None:
    try:
        data = json.loads(_STORE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = {}
    data[_signature(headers)] = mapping
    try:
        _STORE.parent.mkdir(parents=True, exist_ok=True)
        _STORE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass   # ne pas empêcher l'analyse si le dossier est en lecture seule


def _read_csv_text(path: Path) -> str:
    # Les CSV exportés par Excel sous Windows sont souvent en cp1252, pas en UTF-8
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Encodage du fichier non reconnu.")


def _delimiter(text: str) -> str:
    first = text.splitlines()[0] if text else ""
    return ";" if first.count(";") > first.count(",") else ","


def read_headers(path: str | Path) -> list[str]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = _read_csv_text(path)
        if not text.strip():
            raise ValueError("Le fichier est vide.")
        first = text.splitlines()[0]
        return next(csv.reader([first], delimiter=_delimiter(text)), [])
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            if ws is None:
                raise ValueError("Classeur sans feuille active.")
            first = next(ws.iter_rows(values_only=True), None)
            if not first:
                raise ValueError("Le fichier est vide.")
            return [str(h).strip() if h is not None else "" for h in first]
        finally:
            wb.close()
    raise ValueError(f"Format non supporté : {suffix} (utiliser .csv ou .xlsx)")


def read_rows(path: str | Path) -> list[dict]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = _read_csv_text(path)
        return list(csv.DictReader(text.splitlines(), delimiter=_delimiter(text)))
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            if ws is None:
                raise ValueError("Classeur sans feuille active.")
            it = ws.iter_rows(values_only=True)
            headers = [str(h).strip() if h is not None else "" for h in next(it, [])]
            return [dict(zip(headers, r)) for r in it if any(c is not None for c in r)]
        finally:
            wb.close()
    raise ValueError(f"Format non supporté : {suffix} (utiliser .csv ou .xlsx)")


SHEET_KEYS = ("Feuille", "Ligne feuille")


@dataclass
class Sheet:
    name: str
    headers: list[str]
    rows: list[dict]
    numbers: list[int]          # numéro de ligne dans la feuille Excel, pour chaque ligne de données


def _read_xlsx_sheets(path: Path) -> list[Sheet]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[Sheet] = []
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":          # feuilles masquées (listes, calculs) : ignorées
                continue
            headers: list[str] = []
            rows: list[dict] = []
            numbers: list[int] = []
            for number, values in enumerate(ws.iter_rows(values_only=True), start=1):
                if not headers:
                    if any(v is not None and str(v).strip() for v in values):
                        headers = [str(v).strip() if v is not None else "" for v in values]
                    continue
                if any(v is not None for v in values):
                    rows.append(dict(zip(headers, values)))
                    numbers.append(number)
            if rows:
                sheets.append(Sheet(ws.title, headers, rows, numbers))
        return sheets
    finally:
        wb.close()


def read_workbook(path: str | Path) -> list[Sheet]:
    """Les feuilles de données d'un fichier : un classeur Excel peut en avoir plusieurs, un CSV une seule."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = _read_csv_text(path)
        rows = list(csv.DictReader(text.splitlines(), delimiter=_delimiter(text)))
        if not rows:
            return []
        return [Sheet(path.stem, list(rows[0].keys()), rows, list(range(2, len(rows) + 2)))]
    if suffix == ".xlsx":
        return _read_xlsx_sheets(path)
    raise ValueError(f"Format non supporté : {suffix} (utiliser .csv ou .xlsx)")


def looks_like_registrations(sheet: Sheet) -> bool:
    """Vrai si la feuille a des colonnes de nom et de prénom : une liste de personnes."""
    mapping = guess_mapping(sheet.headers)
    return "nom" in mapping and "prenom" in mapping


def combine_sheets(sheets: list[Sheet], label_rows: bool) -> list[dict]:
    """Met bout à bout les lignes des feuilles.

    Les colonnes de même nom (à la casse, aux accents et aux espaces près) sont réunies sous la
    même écriture : la première rencontrée. `label_rows` ajoute la feuille et la ligne d'origine."""
    canonical: dict[str, str] = {}
    for sheet in sheets:
        for header in sheet.headers:
            canonical.setdefault(normalize_name(header), header)
    combined: list[dict] = []
    for sheet in sheets:
        rename = {h: canonical[normalize_name(h)] for h in sheet.headers}
        for row, number in zip(sheet.rows, sheet.numbers):
            merged = {rename.get(k, k): v for k, v in row.items()}
            if label_rows:
                merged["Feuille"] = sheet.name
                merged["Ligne feuille"] = number
            combined.append(merged)
    return combined


def sheet_problems(sheets: list[Sheet], mapping: dict[str, str]) -> tuple[list[str], list[str]]:
    """Colonnes choisies absentes de certaines feuilles : (erreurs bloquantes, avertissements)."""
    errors: list[str] = []
    warnings: list[str] = []
    for sheet in sheets:
        present = {normalize_name(h) for h in sheet.headers}
        for field, column in mapping.items():
            if normalize_name(column) in present:
                continue
            message = f"« {sheet.name} » : colonne « {column} » introuvable"
            (errors if field in ("nom", "prenom") else warnings).append(message)
    return errors, warnings
