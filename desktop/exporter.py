from openpyxl import Workbook

from deduplication import DedupResult, MatchResult, NormalizedRecord
from deduplication.matching.comparators import compare_email, phone_distance
from deduplication.normalization import normalize_name

EMAIL_SIMILAR = 0.85


def _cell(value):
    # Une valeur commençant par "=" serait exécutée comme une formule par Excel
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def _raw(rec: NormalizedRecord, mapping: dict[str, str], key: str) -> str | None:
    col = mapping.get(key)
    value = rec.raw.get(col) if col else None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = "" if value is None else str(value).strip()
    return text or None


def _collect_emails(recs: list[NormalizedRecord]) -> list[str]:
    out: list[str] = []
    for r in recs:
        # une simple faute de frappe (gmial.com) n'est pas un nouvel email
        if r.email and all((compare_email(r.email, e) or 0) < EMAIL_SIMILAR for e in out):
            out.append(r.email)
    return out


def _collect_phones(recs: list[NormalizedRecord]) -> list[str]:
    out: list[str] = []
    for r in recs:
        if r.telephone and all(phone_distance(r.telephone, p) == 2 for p in out):
            out.append(r.telephone)
    return out


def export_workbook(
    path: str,
    records: list[NormalizedRecord],
    final: DedupResult,
    pending: list[tuple[int, int, MatchResult]],
    mapping: dict[str, str],
) -> None:
    headers = list(records[0].raw.keys()) if records else []
    wb = Workbook(write_only=True)

    # --- inscriptions : toutes les lignes, toutes les colonnes d'origine
    ws = wb.create_sheet("inscriptions")
    ws.append(["id_personne", "inscription_en_double", "email_normalise",
               "telephone_normalise", *headers])
    seen: set[tuple[int, str]] = set()
    for i, rec in enumerate(records):
        pid = final.person_of[i]
        formation = rec.extra.get("formation")
        duplicate = ""
        if formation:
            key = (pid, normalize_name(formation))
            duplicate = "oui" if key in seen else ""
            seen.add(key)
        ws.append([pid, duplicate, rec.email, rec.telephone,
                   *[_cell(rec.raw.get(h)) for h in headers]])

    # --- personnes : une ligne par personne, avec toutes ses valeurs
    ws = wb.create_sheet("personnes")
    ws.append(["id_personne", "nom", "prenom", "sexe", "emails", "telephones",
               "nb_inscriptions", "formations"])
    for pid in sorted(final.clusters):
        recs = [records[i] for i in final.clusters[pid]]
        formations = list(dict.fromkeys(f for r in recs if (f := r.extra.get("formation"))))
        sexe = next((r.extra.get("sexe") for r in recs if r.extra.get("sexe")), None)
        ws.append([
            pid,
            _cell(_raw(recs[0], mapping, "nom")),
            _cell(_raw(recs[0], mapping, "prenom")),
            _cell(sexe),
            " | ".join(_collect_emails(recs)),
            " | ".join(_collect_phones(recs)),
            len(recs),
            _cell(" | ".join(formations)),
        ])

    # --- a_verifier : paires non tranchées
    ws = wb.create_sheet("a_verifier")
    ws.append(["ligne_a", "ligne_b", "id_personne_a", "id_personne_b", "score", "raisons"])
    for i, j, res in pending:
        ws.append([records[i].row, records[j].row, final.person_of[i], final.person_of[j],
                   round(res.score, 2), " ; ".join(res.reasons)])

    # --- corrections : normalisations et problèmes de données
    ws = wb.create_sheet("corrections")
    ws.append(["ligne", "champ", "type", "valeur_originale", "valeur_normalisee", "detail"])
    for rec in records:
        for c in rec.corrections:
            if c.normalized is not None and c.original.strip().casefold() == c.normalized:
                continue   # simple changement de casse ou d'espaces : pas la peine
            ws.append([c.row, c.field, "correction", _cell(c.original), _cell(c.normalized), None])
        for issue in rec.issues:
            ws.append([issue.row, issue.field, "probleme", None, None, _cell(issue.message)])

    wb.save(path)