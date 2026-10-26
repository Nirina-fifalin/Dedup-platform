from openpyxl import Workbook

import file_reader


def make_workbook(path):
    wb = Workbook()
    ws = wb.active
    assert ws is not None            # un classeur neuf a toujours une feuille
    ws.title = "2024"
    ws.append(["Nom", "Prénom", "Email"])
    ws.append(["Rakoto", "Jean", "j@x.mg"])
    ws.append(["Rasoa", "Marie", "m@x.mg"])
    ws2 = wb.create_sheet("2025")
    ws2.append(["Nom", "Prénom", "Email"])
    ws2.append(["Andria", "Tiana", "t@x.mg"])
    wb.create_sheet("Vide")                       # sans données : ignorée
    hidden = wb.create_sheet("Listes")
    hidden.append(["Valeur"])
    hidden.append(["x"])
    hidden.sheet_state = "hidden"                 # masquée : ignorée
    wb.active = 1                                 # la feuille active n'est PAS la première
    wb.save(path)


def test_all_visible_sheets_are_read_not_only_the_active_one(tmp_path):
    path = tmp_path / "classeur.xlsx"
    make_workbook(path)
    sheets = file_reader.read_workbook(path)
    assert [(s.name, len(s.rows)) for s in sheets] == [("2024", 2), ("2025", 1)]
    assert sheets[0].numbers == [2, 3]


def test_combine_adds_sheet_and_line_for_multi_sheet_workbooks(tmp_path):
    path = tmp_path / "classeur.xlsx"
    make_workbook(path)
    sheets = file_reader.read_workbook(path)
    rows = file_reader.combine_sheets(sheets, label_rows=True)
    assert len(rows) == 3
    assert (rows[0]["Feuille"], rows[0]["Ligne feuille"]) == ("2024", 2)
    assert (rows[2]["Feuille"], rows[2]["Ligne feuille"]) == ("2025", 2)
    assert "Feuille" not in file_reader.combine_sheets(sheets, label_rows=False)[0]


def test_columns_with_different_spelling_are_unified_and_missing_ones_reported():
    a = file_reader.Sheet("2025", ["Nom", "Prénom", "E-mail"],
                          [{"Nom": "Rakoto", "Prénom": "Jean", "E-mail": "j@x.mg"}], [2])
    b = file_reader.Sheet("2026", ["nom", "prenom", "e-mail", "LinkedIn"],
                          [{"nom": "Rasoa", "prenom": "Marie", "e-mail": "m@x.mg", "LinkedIn": "in/m"}], [2])
    rows = file_reader.combine_sheets([a, b], label_rows=False)
    assert rows[1] == {"Nom": "Rasoa", "Prénom": "Marie", "E-mail": "m@x.mg", "LinkedIn": "in/m"}

    mapping = {"nom": "Nom", "prenom": "Prénom", "email": "E-mail"}
    assert file_reader.sheet_problems([a, b], mapping) == ([], [])

    c = file_reader.Sheet("2027", ["Nom", "E-mail"], [{"Nom": "x", "E-mail": "y"}], [2])
    errors, warnings = file_reader.sheet_problems([a, c], mapping)
    assert errors == ["« 2027 » : colonne « Prénom » introuvable"]
    assert warnings == []


def test_registration_sheets_are_recognised_by_their_name_and_first_name_columns():
    people = file_reader.Sheet("Learners", ["Name", "Last name", "e-mail"], [{}], [2])
    other = file_reader.Sheet("Startups Database", ["Startup", "Sector"], [{}], [2])
    assert file_reader.looks_like_registrations(people)
    assert not file_reader.looks_like_registrations(other)


def test_csv_is_a_single_sheet(tmp_path):
    path = tmp_path / "f.csv"
    path.write_text("Nom;Prénom\nRakoto;Jean\n", encoding="utf-8")
    sheets = file_reader.read_workbook(path)
    assert len(sheets) == 1
    assert sheets[0].rows == [{"Nom": "Rakoto", "Prénom": "Jean"}]


def test_duplicate_column_titles_are_kept_apart(tmp_path):
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Learners"
    ws.append(["Name", "Place", "Training", "Place"])
    ws.append(["Rakoto", "Coding School", "Excel", "Fianarantsoa"])
    path = tmp_path / "dup.xlsx"
    wb.save(path)

    sheet = file_reader.read_workbook(path)[0]
    assert sheet.headers == ["Name", "Place", "Training", "Place (2)"]
    assert sheet.rows[0]["Place"] == "Coding School"          # la 1re colonne n'est plus écrasée
    assert sheet.rows[0]["Place (2)"] == "Fianarantsoa"