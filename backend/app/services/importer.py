import csv
import hashlib
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from deduplication import ColumnMapping, KnownPerson, NormalizedRecord, RecordNormalizer
from deduplication.matching import Classification, Matcher, MatchResult
from deduplication.matching.comparators import compare_email, phone_distance
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session, selectinload

from ..models import (
    Formation, MatchResultRecord, Person, PersonEmail, PersonPhone,
    Registration, SourceFile, ImportLogEntry,
)

_ORDER = {Classification.CERTAIN: 2, Classification.PROBABLE: 1, Classification.NONE: 0}


class AlreadyImportedError(Exception):
    pass


@dataclass
class ImportReport:
    source_file_id: int
    rows: int = 0
    new_persons: int = 0
    matched_existing: int = 0     # inscriptions rattachées à une personne connue
    to_review: int = 0            # cas probables à valider
    skipped: int = 0              # lignes sans nom/prénom
    issues: Counter = field(default_factory=Counter)

    def __str__(self) -> str:
        lines = [
            "─" * 40,
            "RÉSULTAT DU TRAITEMENT",
            "─" * 40,
            f"Inscriptions traitées        {self.rows:>6}",
            f"Personnes déjà connues       {self.matched_existing:>6}",
            f"Nouvelles personnes          {self.new_persons:>6}",
            f"Correspondances à vérifier   {self.to_review:>6}",
            f"Lignes ignorées (sans nom)   {self.skipped:>6}",
            "─" * 40,
        ]
        if self.issues:
            lines.append("Problèmes de données : " + ", ".join(f"{k}={v}" for k, v in self.issues.items()))
        return "\n".join(lines)

    def to_dict(self) -> dict:
        return {
            "source_file_id": self.source_file_id, "rows": self.rows,
            "new_persons": self.new_persons, "matched_existing": self.matched_existing,
            "to_review": self.to_review, "skipped": self.skipped,
            "issues": dict(self.issues),
        }


# ── Lecture des fichiers ────────────────────────────────────────────────────

def _detect_delimiter(path: Path) -> str:
    with path.open(encoding="utf-8-sig") as f:
        head = f.readline()
    return ";" if head.count(";") > head.count(",") else ","


def read_rows(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f, delimiter=_detect_delimiter(path)))
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            if ws is None:
                raise ValueError(f"{path.name} : classeur sans feuille active")
            it = ws.iter_rows(values_only=True)
            headers = [str(h).strip() if h is not None else "" for h in next(it)]
            return [dict(zip(headers, r)) for r in it if any(c is not None for c in r)]
        finally:
            wb.close()
    raise ValueError(f"Format non supporté : {suffix} (utiliser .csv ou .xlsx)")


# ── Aides ───────────────────────────────────────────────────────────────────

def _raw(row: dict, resolved: dict[str, str], name: str) -> str | None:
    col = resolved.get(name)
    value = row.get(col) if col else None
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _jsonable(row: dict) -> dict:
    return {
        str(k): (v if isinstance(v, (str, int, float, bool, type(None))) else str(v))
        for k, v in row.items()
    }


def _to_known(p: Person) -> KnownPerson:
    return KnownPerson(
        id=p.id, nom=p.nom_norm, prenom=p.prenom_norm,
        emails=[e.email for e in p.emails], phones=[x.phone for x in p.phones],
        extra={"sexe": p.sexe},
    )


def _find_candidates(session: Session, rec: NormalizedRecord) -> list[Person]:
    """Blocking : email exact, téléphone exact, ou nom ET prénom *similaires* (trigrammes).

    Permissif par conception : le Matcher tranche ensuite. Le nom seul ne suffit
    pas à devenir candidat (nom ET prénom doivent être proches), dans les deux
    ordres pour couvrir les colonnes inversées.
    """
    conds = []
    if rec.email:
        conds.append(Person.emails.any(PersonEmail.email == rec.email))
    if rec.telephone:
        conds.append(Person.phones.any(PersonPhone.phone == rec.telephone))
    conds.append(and_(
        Person.nom_norm.op("%")(rec.nom), Person.prenom_norm.op("%")(rec.prenom)
    ))
    conds.append(and_(
        Person.nom_norm.op("%")(rec.prenom), Person.prenom_norm.op("%")(rec.nom)
    ))

    closeness = func.greatest(
        func.similarity(Person.nom_norm, rec.nom) + func.similarity(Person.prenom_norm, rec.prenom),
        func.similarity(Person.nom_norm, rec.prenom) + func.similarity(Person.prenom_norm, rec.nom),
    )
    stmt = (
        select(Person)
        .where(Person.merged_into_id.is_(None), or_(*conds))
        .options(selectinload(Person.emails), selectinload(Person.phones))
        .order_by(closeness.desc())
        .limit(50)
    )
    return list(session.scalars(stmt))

def _rank_candidates(matcher: Matcher, rec: NormalizedRecord, candidates: list[Person]):
    scored: list[tuple[MatchResult, Person]] = [
        (matcher.compare_to_person(rec, _to_known(p)), p) for p in candidates
    ]
    scored.sort(key=lambda t: (_ORDER[t[0].classification], t[0].score), reverse=True)
    return scored


def _formation(session: Session, cache: dict, rec: NormalizedRecord) -> Formation | None:
    name = rec.extra.get("formation")
    if not name:
        return None
    annee = rec.extra.get("annee")
    if annee:
        name = f"{name} {annee.removesuffix('.0')}"
    if name not in cache:
        f = session.scalar(select(Formation).where(Formation.nom == name))
        if f is None:
            f = Formation(nom=name)
            session.add(f)
            session.flush()
        cache[name] = f
    return cache[name]


def _parse_date(value: str | None) -> date | None:
    try:
        return date.fromisoformat(value[:10]) if value else None
    except ValueError:
        return None


def _add_contacts(person: Person, rec: NormalizedRecord, reg: Registration, matcher: Matcher) -> None:
    """Ajoute l'email/le téléphone de l'inscription aux valeurs connues de la personne.

    On n'ajoute PAS une valeur qui n'est qu'une faute de frappe d'une valeur déjà
    connue (ex. gmial.com) : l'inscription la conserve de toute façon.
    """
    if rec.email:
        similar = any(
            (compare_email(rec.email, e.email) or 0) >= matcher.config.email_similar_threshold
            for e in person.emails
        )
        if not similar:
            person.emails.append(PersonEmail(
                email=rec.email, is_primary=not person.emails,
                first_seen_registration_id=reg.id,
            ))
    if rec.telephone:
        if all(phone_distance(rec.telephone, p.phone) == 2 for p in person.phones):
            person.phones.append(PersonPhone(
                phone=rec.telephone, is_primary=not person.phones,
                first_seen_registration_id=reg.id,
            ))
    if person.sexe is None and rec.extra.get("sexe"):
        person.sexe = rec.extra["sexe"]


def _is_cosmetic(original: str, normalized: str | None) -> bool:
    """Majuscules/espaces seulement : pas la peine de le journaliser."""
    return normalized is not None and original.strip().casefold() == normalized

# ── Import ──────────────────────────────────────────────────────────────────

def import_file(
    session: Session,
    path: str | Path,
    mapping: ColumnMapping,
    matcher: Matcher | None = None,
    commit: bool = True,
) -> ImportReport:
    path = Path(path)
    matcher = matcher or Matcher()

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if session.scalar(
        select(SourceFile.id).where(SourceFile.file_hash == digest, SourceFile.status == "done")
    ):
        raise AlreadyImportedError(f"{path.name} a déjà été importé (même contenu).")

    rows = read_rows(path)
    normalizer = RecordNormalizer(mapping)
    resolved = normalizer.resolve_columns(rows[0].keys()) if rows else {}

    source = SourceFile(filename=path.name, file_hash=digest, row_count=len(rows), status="processing")
    session.add(source)
    session.flush()
    report = ImportReport(source_file_id=source.id, rows=len(rows))

    formations: dict[str, Formation] = {}

    try:
        for row, rec in zip(rows, normalizer.normalize_rows(rows)):
            for corr in rec.corrections:
                if _is_cosmetic(corr.original, corr.normalized):
                    continue
                session.add(ImportLogEntry(
                    source_file_id=source.id, source_row=corr.row, field=corr.field,
                    kind="correction", original=corr.original, normalized=corr.normalized,
                ))
            for issue in rec.issues:
                report.issues[issue.code] += 1
                session.add(ImportLogEntry(
                    source_file_id=source.id, source_row=issue.row, field=issue.field,
                    kind="issue", code=issue.code, message=issue.message,
                ))
            if not rec.nom or not rec.prenom:
                report.skipped += 1
                continue

            ranked = _rank_candidates(matcher, rec, _find_candidates(session, rec))
            top_res, top_person = ranked[0] if ranked else (None, None)

            if (
                top_res is not None
                and top_person is not None
                and top_res.classification == Classification.CERTAIN
            ):
                certain = True
                person = top_person
                report.matched_existing += 1
            else:
                certain = False
                person = Person(
                    nom=_raw(row, resolved, "nom") or rec.nom,
                    prenom=_raw(row, resolved, "prenom") or rec.prenom,
                    nom_norm=rec.nom, prenom_norm=rec.prenom,
                    sexe=rec.extra.get("sexe"),
                )
                session.add(person)
                session.flush()
                report.new_persons += 1

            formation = _formation(session, formations, rec)
            reg = Registration(
                person_id=person.id,
                formation_id=formation.id if formation else None,
                source_file_id=source.id, source_row=rec.row,
                date_inscription=_parse_date(rec.extra.get("date_inscription")),
                nom_submitted=_raw(row, resolved, "nom"),
                prenom_submitted=_raw(row, resolved, "prenom"),
                email_submitted=_raw(row, resolved, "email"),
                phone_submitted=_raw(row, resolved, "telephone"),
                raw=_jsonable(row),
            )
            session.add(reg)
            session.flush()

            _add_contacts(person, rec, reg, matcher)

            if (
                not certain
                and top_res is not None
                and top_person is not None
                and top_res.classification == Classification.PROBABLE
            ):
                session.add(MatchResultRecord(
                    registration_id=reg.id, person_a_id=person.id, person_b_id=top_person.id,
                    score=top_res.score, classification=top_res.classification.value,
                    reasons=top_res.reasons, conflicts=top_res.conflicts,
                ))
                report.to_review += 1

        source.report = report.to_dict()
        source.status = "done"
        session.flush()
        if commit:
            session.commit()
    except Exception:
        session.rollback()
        raise
    return report