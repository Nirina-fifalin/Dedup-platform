from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from ..models import (
    Formation, ImportLogEntry, MatchResultRecord, Person, Registration, SourceFile,
)

BATCH = 1000


def _cell(value):
    """Une chaîne commençant par '=' serait interprétée comme une formule par Excel."""
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def _row(values) -> list:
    return [_cell(v) for v in values]


def build_export(session: Session, source_file_id: int) -> bytes:
    if session.get(SourceFile, source_file_id) is None:
        raise LookupError("Fichier source introuvable")

    wb = Workbook(write_only=True)
    file_persons = (
        select(Registration.person_id)
        .where(Registration.source_file_id == source_file_id).distinct()
    )

    # --- personnes
    counts = dict(session.execute(
        select(Registration.person_id, func.count())
        .where(Registration.person_id.in_(file_persons))
        .group_by(Registration.person_id)
    ).all())
    ws = wb.create_sheet("personnes")
    ws.append(["person_id", "nom", "prenom", "sexe", "emails", "telephones", "nb_inscriptions_total"])
    persons = session.scalars(
        select(Person).where(Person.id.in_(file_persons)).order_by(Person.id)
        .options(selectinload(Person.emails), selectinload(Person.phones))
        .execution_options(yield_per=BATCH)
    )
    for p in persons:
        emails = " | ".join(e.email for e in sorted(p.emails, key=lambda e: not e.is_primary))
        phones = " | ".join(x.phone for x in sorted(p.phones, key=lambda x: not x.is_primary))
        ws.append(_row([p.id, p.nom, p.prenom, p.sexe, emails, phones, counts.get(p.id, 0)]))

    # --- inscriptions
    ws = wb.create_sheet("inscriptions")
    ws.append(["registration_id", "person_id", "formation", "date_inscription", "ligne_source",
               "nom_saisi", "prenom_saisi", "email_saisi", "telephone_saisi"])
    regs = session.execute(
        select(Registration, Formation.nom)
        .outerjoin(Formation, Formation.id == Registration.formation_id)
        .where(Registration.source_file_id == source_file_id)
        .order_by(Registration.source_row, Registration.id)
        .execution_options(yield_per=BATCH)
    )
    for r, formation in regs:
        ws.append(_row([r.id, r.person_id, formation, r.date_inscription, r.source_row,
                        r.nom_submitted, r.prenom_submitted, r.email_submitted, r.phone_submitted]))

    # --- doublons et a_verifier
    reg_ids = select(Registration.id).where(Registration.source_file_id == source_file_id)
    cases = session.scalars(
        select(MatchResultRecord).where(MatchResultRecord.registration_id.in_(reg_ids))
        .order_by(MatchResultRecord.score.desc(), MatchResultRecord.id)
    ).all()
    headers = ["case_id", "person_a", "person_b", "score", "classification", "statut", "raisons", "conflits"]

    def case_row(c: MatchResultRecord) -> list:
        return _row([c.id, c.person_a_id, c.person_b_id, c.score, c.classification, c.status,
                     " ; ".join(c.reasons or []), ", ".join(c.conflicts or [])])

    ws = wb.create_sheet("doublons")
    ws.append(headers)
    for c in cases:
        ws.append(case_row(c))

    ws = wb.create_sheet("a_verifier")
    ws.append(headers)
    for c in cases:
        if c.status in ("pending", "postponed"):
            ws.append(case_row(c))

    # --- corrections
    ws = wb.create_sheet("corrections")
    ws.append(["ligne", "champ", "type", "valeur_originale", "valeur_normalisee", "code", "message"])
    log = session.scalars(
        select(ImportLogEntry).where(ImportLogEntry.source_file_id == source_file_id)
        .order_by(ImportLogEntry.source_row, ImportLogEntry.id)
        .execution_options(yield_per=BATCH)
    )
    for e in log:
        kind = "correction" if e.kind == "correction" else "probleme"
        ws.append(_row([e.source_row, e.field, kind, e.original, e.normalized, e.code, e.message]))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()