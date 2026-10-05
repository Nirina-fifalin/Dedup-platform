import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.db import engine
from app.deps import get_session
from app.main import app


@pytest.fixture
def client():
    conn = engine.connect()
    outer = conn.begin()
    session = Session(bind=conn, join_transaction_mode="create_savepoint", expire_on_commit=False)
    app.dependency_overrides[get_session] = lambda: session
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()
        session.close()
        outer.rollback()
        conn.close()


CSV_SAME = (
    "nom,prenom,email,telephone\n"
    "Apitest,Zeta,zeta.apitest@example.test,0349922201\n"
    "Apitest,Zeta,zeta.apitest@example.test,0349922201\n"
)
CSV_CONFLICT = (
    "nom,prenom,email,telephone\n"
    "Apitest,Eta,eta.apitest@example.test,0349922202\n"
    "Apitest,Eta,eta.apitest@example.test,0329922299\n"
)


def upload(client, content, name="api.csv"):
    return client.post("/api/v1/files", files={"file": (name, content, "text/csv")})


def test_upload_reports_and_rejects_duplicate_file(client):
    r = upload(client, CSV_SAME)
    assert r.status_code == 201
    body = r.json()
    assert (body["rows"], body["new_persons"], body["matched_existing"]) == (2, 1, 1)

    assert upload(client, CSV_SAME).status_code == 409

    info = client.get(f"/api/v1/files/{body['source_file_id']}").json()
    assert info["status"] == "done"
    assert info["report"]["rows"] == 2


def test_rejects_unsupported_format(client):
    r = client.post("/api/v1/files", files={"file": ("x.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 415


def test_rejects_missing_required_columns(client):
    r = upload(client, "nom,email\nApitest,a@example.test\n", name="bad.csv")
    assert r.status_code == 422


def test_review_queue_and_decision(client):
    assert upload(client, CSV_CONFLICT, name="conflict.csv").status_code == 201

    cases = client.get("/api/v1/review?limit=200").json()
    mine = next(c for c in cases if c["person_a"]["nom"] == "Apitest")
    assert "telephone_conflict" in mine["conflicts"]
    assert mine["person_b"]["registrations"] == 1

    r = client.post(
        f"/api/v1/review/{mine['id']}/decision",
        json={"decision": "kept_separate", "decided_by": "test"},
    )
    assert r.status_code == 200
    assert r.json()["status"] == "kept_separate"

    again = client.post(f"/api/v1/review/{mine['id']}/decision", json={"decision": "postponed"})
    assert again.status_code == 409


CSV_MERGE = (
    "nom,prenom,email,telephone\n"
    "Apitest,Theta,theta.apitest@example.test,0349922203\n"
    "Apitest,Theta,theta.apitest@example.test,0329922298\n"
)


def test_merge_decision_then_undo_reopens_the_case(client):
    assert upload(client, CSV_MERGE, name="merge.csv").status_code == 201
    case = next(c for c in client.get("/api/v1/review?limit=200").json()
                if c["person_a"]["prenom"] == "Theta")

    r = client.post(f"/api/v1/review/{case['id']}/decision",
                    json={"decision": "merged", "decided_by": "test"})
    assert r.status_code == 200
    assert r.json()["status"] == "merged"
    assert r.json()["person_a"]["registrations"] == 0
    assert r.json()["person_b"]["registrations"] == 2
    assert len(r.json()["person_b"]["phones"]) == 2
    assert len(r.json()["person_b"]["emails"]) == 1

    merge = next(m for m in client.get("/api/v1/merges").json() if m["case_id"] == case["id"])
    assert merge["registrations_moved"] == 1

    undo = client.post(f"/api/v1/merges/{merge['id']}/undo", json={"undone_by": "test"})
    assert undo.status_code == 200
    assert undo.json()["undone_at"] is not None

    back = client.get(f"/api/v1/review/{case['id']}").json()
    assert back["status"] == "pending"
    assert back["person_a"]["registrations"] == 1
    assert back["person_a"]["emails"] == ["theta.apitest@example.test"]

    assert client.post(f"/api/v1/merges/{merge['id']}/undo", json={}).status_code == 409

from io import BytesIO

from openpyxl import load_workbook

CSV_EXPORT = (
    "nom,prenom,email,telephone\n"
    "EXPORTTEST,Iota,Iota.Exporttest@EXAMPLE.test,034 99 222 04\n"
    "Exporttest,Iota,iota.exporttest@example.test,0329922297\n"
    "=1+1,Kappa,kappa.exporttest@example.test,0349922205\n"
)


def test_export_sheets_corrections_and_formula_guard(client):
    body = upload(client, CSV_EXPORT, name="export.csv").json()
    r = client.get(f"/api/v1/files/{body['source_file_id']}/export")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]

    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["personnes", "inscriptions", "doublons", "a_verifier", "corrections"]
    assert wb["personnes"].max_row == 1 + 3
    assert wb["inscriptions"].max_row == 1 + 3
    assert wb["doublons"].max_row == 1 + 1        # même email, téléphones différents
    assert wb["a_verifier"].max_row == 1 + 1

    # Valeur saisie conservée, mais neutralisée pour Excel
    inscriptions = list(wb["inscriptions"].iter_rows(min_row=2, values_only=True))
    assert inscriptions[2][5] == "'=1+1"

    corrections = list(wb["corrections"].iter_rows(min_row=2, values_only=True))
    assert any(c[0] == 2 and c[1] == "telephone" and c[3] == "034 99 222 04" for c in corrections)
    # les simples changements de casse ne polluent pas la feuille
    assert not any(c[0] == 2 and c[1] in ("nom", "email") and c[2] == "correction" for c in corrections)